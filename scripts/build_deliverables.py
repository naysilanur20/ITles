"""Generate reviewable deliverables from versioned evidence, not the input workbook."""

from __future__ import annotations

import json
import hashlib
from datetime import datetime, timezone
from pathlib import Path
import re
import textwrap
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
DATE = "2026-09-15"
LABELS = {
    "id": "ID", "title": "Документ", "publisher": "Издатель", "date": "Дата документа",
    "url": "Ссылка", "claim": "Проверяемое утверждение", "status": "Статус подтверждения",
    "model": "Модель", "head": "Головка", "computer": "Бортовая система",
    "confirmed": "Что подтверждено", "access": "Получение данных", "restrictions": "Ограничения",
    "extra": "Что потребуется", "sources": "Источники / ID", "metric": "Показатель",
    "purpose": "Назначение", "source": "Источник данных", "location": "Узел / расположение",
    "unit": "Единицы", "interface": "Способ получения", "frequency": "Частота обновления",
    "accuracy": "Точность", "limitations": "Ограничения", "retrofit": "Доработки",
    "option": "Вариант", "hardware": "Оборудование", "validation": "Проверка",
    "region": "Регион", "enterprise": "Предприятие", "evidence": "Основание / результат",
    "volume": "Объём / охват", "verdict": "Вывод проверки", "reason": "Обоснование",
    "question": "Открытый вопрос", "why": "Зачем выяснить", "verification": "Как проверить",
    "name": "Название", "inputs": "Исходные данные", "formula": "Формула",
    "result": "Результат", "meaning": "Интерпретация", "limitation": "Граница применимости",
    "field": "Поле ИТлес", "target": "Целевое поле / сущность", "transformation": "Преобразование",
    "test": "Испытание", "equipment": "Оборудование", "steps": "Действия",
    "expected": "Ожидаемый результат", "pass_criterion": "Критерий приёмки",
}


def load_evidence():
    files = [ROOT / "research/hardware.json", ROOT / "research/volume-1c.json"]
    for path in files:
        if not path.is_file():
            raise SystemExit(f"Evidence file not ready: {path.relative_to(ROOT)}")
    return [json.loads(path.read_text()) for path in files]


def value_text(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value) if value is not None else "Не установлено"


def make_workbook(evidence):
    book = Workbook()
    book.remove(book.active)
    sheets = [
        ("Как читать", [{"name": "Область исследования", "result": "Проверка документальных источников. Не обследование национального парка и не испытание реальной техники.", "status": "Без полевых данных"},
                         {"name": "Исходная книга", "result": "harvesters_rf.xlsx сохранена без изменения в приложении к заданию. Её сведения не приняты автоматически за факты.", "status": "Аудит на отдельном листе"},
                         {"name": "Проценты и регионы", "result": "Доли рынка и оснащённости не рассчитаны: нет репрезентативного реестра и проверенного знаменателя. Региональные строки — примеры с ограниченным охватом, не рейтинг.", "status": "Национальная оценка отсутствует"},
                         {"name": "Комплектация", "result": "Возможность в документации производителя не доказывает наличие, исправность и доступность данных на конкретном серийном номере.", "status": "Требуется проверка конфигурации"},
                         {"name": "Персональные данные", "result": "Контакты людей и данные операторов не включены. Реальная геолокация вместе со сменным графиком может стать персональной.", "status": "Требуется согласование до пилота"},
                         {"name": "Дата проверки", "result": DATE, "status": "Даты публикаций указаны отдельно"}]),
        ("Уровни доказательности", [
            {"name": "Подтверждено документацией", "result": "Возможность изделия описана в указанном первичном источнике. Не означает оснащённость, исправность или доступность интерфейса конкретной машины.", "status": "Исследовано"},
            {"name": "Спроектировано", "result": "Архитектура OEM-адаптера, схема сверки, программа пилота и обмен с согласуемой конфигурацией 1С.", "status": "Не подключено"},
            {"name": "Реализовано", "result": "Код нормализованного приёмника, очереди, БД, расчётов и интерфейса. Сам факт наличия кода не доказывает правильность.", "status": "См. реестр проверки в отчёте"},
            {"name": "Проверено синтетически", "result": "Автоматические тесты специально составленных событий. Фактические команды и результаты фиксируются отдельно в отчёте.", "status": "Не метрологическое испытание"},
            {"name": "Проверено на реальных OEM-файлах", "result": "Эталонные HPR/StanForD-файлы и сырая телеметрия заказчика не предоставлены.", "status": "Не выполнено"},
            {"name": "Проверено в поле", "result": "Доступа к физическому харвестеру, датчикам и приёмке древесины не было.", "status": "Не выполнено"},
        ]),
    ]
    mapping = [
        ("Конфигурации", "configurations"), ("Параметры и датчики", "metrics"),
        ("Подключение", "connections"), ("Регионы", "regions"),
        ("Проверка исходного Excel", "audit"), ("Пробелы и вопросы", "open_questions"),
        ("Расчёты объёма", "volume_checks"), ("Обмен с 1С — проект", "integration"),
        ("Полевые испытания", "field_tests"), ("Источники", "sources"),
    ]
    for title, key in mapping:
        rows = [row for item in evidence for row in item.get(key, [])]
        if not rows:
            raise ValueError(f"Missing evidence section: {key}")
        sheets.append((title, rows))
    enterprises = [row for item in evidence for row in item.get("regions", []) if row.get("enterprise")]
    sheets.insert(5, ("Предприятия", enterprises))
    for index, (title, rows) in enumerate(sheets, 1):
        ws = book.create_sheet(title)
        keys = list(dict.fromkeys(key for row in rows for key in row))
        ws.append([f"ИТлес / {title} / проверка {DATE}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(keys))
        ws.cell(1, 1).font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
        ws.cell(1, 1).fill = PatternFill("solid", fgColor="173D34")
        ws.row_dimensions[1].height = 32
        ws.append(["Документальная возможность ≠ подключённая и исправная машина. Фильтруйте по статусу и проверяйте источник."])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(keys))
        ws.cell(2, 1).font = Font(name="Calibri", italic=True, color="5C6963", size=10)
        ws.row_dimensions[2].height = 25
        ws.append([LABELS.get(key, key) for key in keys])
        for row in rows:
            cells = []
            for key in keys:
                value = value_text(row.get(key, "Не установлено"))
                if value.startswith(("=", "+", "-", "@")):
                    value = "'" + value
                cells.append(value)
            ws.append(cells)
        tab = Table(displayName=f"Evidence{index}", ref=f"A3:{get_column_letter(len(keys))}{ws.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
        for col, key in enumerate(keys, 1):
            width = 44 if key not in ("id", "unit", "date", "status", "sources") else {"id": 14, "unit": 18, "date": 22, "status": 32, "sources": 22}[key]
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in ws.iter_rows(min_row=3):
            estimated_lines = 1
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Calibri", size=11, bold=cell.row == 3, color="173D34")
                width = ws.column_dimensions[cell.column_letter].width
                estimated_lines = max(estimated_lines, sum(max(1, len(line) // int(width - 2) + 1) for line in str(cell.value or "").splitlines()))
                if str(cell.value).startswith(("https://", "http://")):
                    cell.hyperlink = cell.value
                    cell.font = Font(name="Calibri", size=11, color="146F8A", underline="single")
            ws.row_dimensions[row[0].row].height = min(400, max(32, estimated_lines * 15 + 10))
        ws.freeze_panes = "C4" if len(keys) > 2 else "A4"
        ws.sheet_view.zoomScale = 85
        ws.print_title_rows = "1:3"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.oddFooter.center.text = "ИТлес · &A · &P / &N"
    book.properties.creator = "ИТлес"
    book.properties.title = "Харвестеры: источники, доступность данных и проверки"
    book.save(OUT / "harvesters_evidence.xlsx")


def plain_inline(token):
    parts, links = [], []
    for child in token.children or []:
        if child.type in ("text", "code_inline"):
            parts.append(child.content)
        elif child.type in ("softbreak", "hardbreak"):
            parts.append(" ")
        elif child.type == "link_open":
            href = child.attrGet("href")
            if href and href.startswith(("https://", "http://")):
                links.append(href)
    text = "".join(parts)
    for link in links:
        if link not in text:
            text += f" ({link})"
    return text


def report_markdown(evidence):
    generated_on = datetime.now(timezone.utc).date().isoformat()
    chunks = ["# ИТлес\n\n## Технический отчёт и программа пилота\n\nСборка отчёта: " + generated_on + ". Проверка документальных источников: " + DATE + ".\n\nИсследование: открытая документация. Программные примеры: синтетические. Реальная машина и 1С заказчика: не подключены.\n"]
    for path in ("docs/architecture.md", "docs/administrator.md", "docs/user-guide.md", "docs/deployment.md", "docs/api-contract-v2.md", "research/hardware.md", "research/volume-1c.md", "docs/privacy.md", "docs/demo-scenario.md", "docs/verification.md", "README.md"):
        file = ROOT / path
        if file.exists():
            chunks.append(file.read_text())
    results = ROOT / "docs/verification-results.json"
    if results.exists():
        recorded = json.loads(results.read_text())
        current = recorded.get("source_fingerprint") == verification_fingerprint()
        notice = "Отпечаток исходников совпадает с проверенным набором." if current else "ВНИМАНИЕ: это исторический прогон; исходники изменились, его результат не подтверждает текущую версию."
        chunks.append("# Фактически выполненные проверки\n\n" + notice + "\n\n```json\n" + results.read_text() + "\n```\n")
    chunks.append("# Единый реестр источников\n\nДата обращения: " + DATE + ". Возможность, описанная в источнике, не равна совместимости конкретной машины.\n")
    for item in evidence:
        for source in item.get("sources", []):
            chunks.append(f"## {source['id']} · {source['title']}\n\nИздатель: {source.get('publisher', 'Не указан')}. Дата: {source.get('date', 'Не указана')}.\n\n{source.get('claim', '')}\n\nСтатус: {source.get('status', 'Документальный источник')}.\n\n{source['url']}\n")
    return "\n\n".join(chunks)


def make_reports(evidence):
    markdown = report_markdown(evidence)
    (OUT / "itles_report.md").write_text(markdown)
    document = Document()
    document.core_properties.author = "ИТлес"
    document.core_properties.title = "Мониторинг харвестеров: доказательная база и пилот"
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    for style in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[style].font.color.rgb = RGBColor.from_string("173D34")
    section.footer.paragraphs[0].text = "ИТлес · Документация и синтетические проверки · Не полевые испытания"

    fonts = Path("/usr/share/fonts/truetype/dejavu")
    regular = fonts / "DejaVuSans.ttf"
    bold = fonts / "DejaVuSans-Bold.ttf"
    mono = fonts / "DejaVuSansMono.ttf"
    if not regular.exists():
        raise SystemExit("Install DejaVu Sans fonts to generate a Cyrillic PDF")
    for name, path in (("ITles", regular), ("ITles-Bold", bold), ("ITles-Mono", mono)):
        pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily("ITles", normal="ITles", bold="ITles-Bold", italic="ITles", boldItalic="ITles-Bold")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("RussianBody", fontName="ITles", fontSize=9, leading=14, spaceAfter=7, textColor=colors.HexColor("#253B32"), alignment=TA_LEFT, wordWrap="CJK")
    headings = {level: ParagraphStyle(f"RussianH{level}", parent=body, fontName="ITles-Bold", fontSize={1: 19, 2: 13, 3: 11}.get(level, 10), leading={1: 25, 2: 18, 3: 15}.get(level, 14), spaceBefore=15, spaceAfter=9, keepWithNext=True) for level in range(1, 7)}
    code = ParagraphStyle("RussianCode", parent=body, fontName="ITles-Mono", fontSize=7.5, leading=11, backColor=colors.HexColor("#F0F3F0"), borderPadding=6)
    story, mode, list_depth, in_table = [], None, 0, False
    table_row, table_rows = [], []
    parser = MarkdownIt("commonmark").enable("table")

    def add_paragraph(text, level=None, bullet=False):
        if not text.strip():
            return
        document.add_paragraph(text, style=f"Heading {min(level, 3)}" if level else ("List Bullet" if bullet else "Normal"))
        text = escape(text).replace("\n", "<br/>")
        text = re.sub(r"https?://[^\s<]+", lambda m: f'<link href="{m.group(0)}" color="#146F8A">{m.group(0)}</link>', text)
        story.append(Paragraph(("• " if bullet else "") + text, headings[level] if level else body))

    for token in parser.parse(markdown):
        if token.type == "heading_open":
            mode = int(token.tag[1:])
        elif token.type == "heading_close":
            mode = None
        elif token.type in ("bullet_list_open", "ordered_list_open"):
            list_depth += 1
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            list_depth -= 1
        elif token.type == "table_open":
            in_table, table_rows = True, []
        elif token.type == "tr_open":
            table_row = []
        elif token.type == "tr_close":
            table_rows.append(table_row)
        elif token.type == "table_close":
            if table_rows:
                headers = table_rows[0]
                table = document.add_table(rows=1, cols=len(headers))
                table.style = "Light Shading Accent 1"
                for i, text in enumerate(headers):
                    table.rows[0].cells[i].text = text
                for row in table_rows[1:]:
                    cells = table.add_row().cells
                    for i, text in enumerate(row[:len(headers)]):
                        cells[i].text = text
                    text = "<br/>".join(f"<b>{escape(headers[i])}:</b> {escape(text)}" for i, text in enumerate(row[:len(headers)]))
                    story.append(Paragraph(text, body))
                    story.append(Spacer(1, 5))
            in_table = False
        elif token.type == "inline":
            text = plain_inline(token)
            if in_table:
                table_row.append(text)
            else:
                add_paragraph(text, mode, bool(list_depth) and mode is None)
        elif token.type in ("fence", "code_block"):
            document.add_paragraph(token.content, style="No Spacing")
            for line in token.content.splitlines():
                for part in textwrap.wrap(line, width=100, replace_whitespace=False) or [" "]:
                    story.append(Paragraph(escape(part).replace(" ", "&#160;"), code))
            story.append(Spacer(1, 8))
        elif token.type == "hr":
            story.append(Spacer(1, 12))
    document.save(OUT / "itles_report.docx")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("ITles", 7)
        canvas.setFillColor(colors.HexColor("#64746B"))
        canvas.drawString(38, 23, "ИТлес · Не подтверждает подключение к реальной машине")
        canvas.drawRightString(A4[0] - 38, 23, str(doc.page))
        canvas.restoreState()

    pdf = SimpleDocTemplate(str(OUT / "itles_report.pdf"), pagesize=A4, rightMargin=38, leftMargin=38, topMargin=36, bottomMargin=40, title="ИТлес — доказательная база и пилот", author="ИТлес")
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


def source_files():
    included_dirs = ("backend", "frontend/src", "frontend/public", "frontend/e2e", "edge", "scripts", "tests", "research", "docs")
    source_extensions = {".py", ".md", ".json", ".geojson", ".ts", ".tsx", ".css", ".svg", ".png", ".txt", ".ttf"}
    files = [ROOT / name for name in ("README.md", "requirements.txt", "requirements.lock", "pyproject.toml", "server.py", ".gitignore", ".env.example", ".hoplite/settings.json", ".hoplite/setup.sh", ".hoplite/run.sh", ".github/workflows/verify.yml", "frontend/package.json", "frontend/package-lock.json", "frontend/index.html", "frontend/tsconfig.json", "frontend/tsconfig.node.json", "frontend/tsconfig.app.json", "frontend/vite.config.ts", "frontend/vitest.config.ts")]
    for directory in included_dirs:
        for path in (ROOT / directory).rglob("*"):
            parts = path.relative_to(ROOT).parts
            if (
                path.suffix in source_extensions
                and not any(part.startswith(".") or part in {"__pycache__", "node_modules", "dist"} for part in parts)
                and ".hoplite-write-" not in path.name
            ):
                files.append(path)
    files.extend(ROOT / name for name in ("DESIGN.md", "HANDOFF.md", "frontend/playwright.config.ts", ".gitattributes", "Dockerfile", ".dockerignore", "deployment/compose.yaml"))
    return [path for path in sorted(set(files)) if path.is_file() and not path.is_symlink() and path.resolve().is_relative_to(ROOT)]


def verification_fingerprint():
    digest = hashlib.sha256()
    for path in source_files():
        if path.relative_to(ROOT).as_posix() == "docs/verification-results.json":
            continue
        digest.update(path.relative_to(ROOT).as_posix().encode() + b"\0" + path.read_bytes() + b"\0")
    return digest.hexdigest()


def make_source_archive():
    with zipfile.ZipFile(OUT / "itles_source.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in source_files():
            archive_source(archive, path, "itles/" + str(path.relative_to(ROOT)))


def make_deployment_archive():
    with zipfile.ZipFile(OUT / "itles_deploy.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in source_files():
            archive_source(archive, path, str(path.relative_to(ROOT)))


def archive_source(archive, path, name):
    entry = zipfile.ZipInfo.from_file(path, name)
    # Do not inherit private sandbox file modes into the deployable sources.
    entry.external_attr = 0o100644 << 16
    entry.compress_type = zipfile.ZIP_DEFLATED
    archive.writestr(entry, path.read_bytes())


def main():
    OUT.mkdir(exist_ok=True)
    evidence = load_evidence()
    make_workbook(evidence)
    make_reports(evidence)
    make_source_archive()
    make_deployment_archive()
    for file in sorted(OUT.iterdir()):
        print(f"{file.name}: {file.stat().st_size} bytes")


if __name__ == "__main__":
    main()
