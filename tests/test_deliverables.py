"""End-to-end checks for the evidence deliverables generated from versioned data."""

import json
from pathlib import Path
import re
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook

from scripts import build_deliverables


def source_records():
    return [
        source
        for filename in ("hardware.json", "volume-1c.json")
        for source in json.loads((build_deliverables.ROOT / "research" / filename).read_text())["sources"]
    ]


@pytest.fixture(scope="module")
def output_dir(tmp_path_factory):
    output = tmp_path_factory.mktemp("deliverables")
    original_output = build_deliverables.OUT
    build_deliverables.OUT = output
    try:
        build_deliverables.main()
        yield output
    finally:
        build_deliverables.OUT = original_output


def test_evidence_source_ids_are_unique_across_all_input_registers():
    ids = [source["id"] for source in source_records()]
    assert ids
    assert len(ids) == len(set(ids)), "source IDs must identify one source across both evidence files"


def test_workbook_has_filterable_tables_and_linked_complete_source_register(output_dir):
    workbook = load_workbook(output_dir / "harvesters_evidence.xlsx", data_only=False)
    expected_sheets = {
        "Как читать", "Конфигурации", "Параметры и датчики", "Подключение",
        "Регионы", "Предприятия", "Проверка исходного Excel", "Пробелы и вопросы",
        "Расчёты объёма", "Обмен с 1С — проект", "Полевые испытания", "Источники",
    }
    assert expected_sheets <= set(workbook.sheetnames)
    for name in expected_sheets:
        sheet = workbook[name]
        assert len(sheet.tables) == 1, f"{name} must remain a filterable evidence table"
        table = next(iter(sheet.tables.values()))
        assert table.ref.startswith("A3:")
        assert table.ref.endswith(str(sheet.max_row))
        assert sheet.freeze_panes in {"A4", "C4"}

    sources_sheet = workbook["Источники"]
    headers = {cell.value: cell.column for cell in sources_sheet[3]}
    assert {"ID", "Ссылка"} <= set(headers)
    source_rows = list(sources_sheet.iter_rows(min_row=4, values_only=False))
    produced = {row[headers["ID"] - 1].value: row for row in source_rows}
    expected = {source["id"]: source for source in source_records()}
    assert set(produced) == set(expected)
    for source_id, source in expected.items():
        url_cell = produced[source_id][headers["Ссылка"] - 1]
        assert url_cell.value == source["url"]
        assert url_cell.hyperlink is not None
        assert url_cell.hyperlink.target == source["url"]


def _docx_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def test_docx_and_pdf_preserve_russian_report_content_and_valid_document_structure(output_dir):
    docx_path = output_dir / "itles_report.docx"
    text = _docx_text(docx_path)
    assert "Технический отчёт и программа пилота" in text
    assert "Сборка отчёта: " in text
    assert "Единый реестр источников" in text
    assert "Реальная машина и 1С заказчика: не подключены." in text
    assert all(source["url"] in text for source in source_records())

    pdf = (output_dir / "itles_report.pdf").read_bytes()
    assert pdf.startswith(b"%PDF-")
    assert re.search(rb"/Type\s*/Page\b", pdf)
    assert b"/ToUnicode" in pdf, "a Cyrillic report needs Unicode character mappings"
    assert b"startxref" in pdf and pdf.rstrip().endswith(b"%%EOF")


def test_source_archive_contains_reviewable_sources_but_not_secrets_or_local_state(output_dir):
    with zipfile.ZipFile(output_dir / "itles_source.zip") as archive:
        names = archive.namelist()

    assert "itles/edge/outbox.py" in names
    assert "itles/scripts/build_deliverables.py" in names
    assert "itles/research/hardware.json" in names
    assert {"itles/DESIGN.md", "itles/docs/demo-scenario.md", "itles/frontend/src/api.ts",
            "itles/frontend/src/assets/fonts/GolosText.ttf", "itles/frontend/src/assets/fonts/GolosText-OFL.txt",
            "itles/frontend/src/assets/fonts/PTSerif-Regular.ttf", "itles/frontend/src/assets/fonts/PTSerif-OFL.txt"} <= set(names)
    assert all(name.startswith("itles/") and not name.startswith("/") for name in names)

    forbidden_parts = {".git", ".local", ".venv", "__pycache__", ".pytest_cache", "node_modules", "data", "deliverables"}
    forbidden_suffixes = (".db", ".sqlite", ".sqlite3", ".pyc", ".pem", ".key")
    for name in names:
        relative = Path(name).relative_to("itles")
        assert not forbidden_parts.intersection(relative.parts), name
        assert not relative.name.endswith(forbidden_suffixes), name
        assert relative.name == ".env.example" or not relative.name.startswith(".env"), name


def test_report_does_not_present_old_verification_as_current(tmp_path, monkeypatch):
    monkeypatch.setattr(build_deliverables, "ROOT", tmp_path)
    (tmp_path / "backend").mkdir()
    (tmp_path / "docs").mkdir()
    source = tmp_path / "backend" / "probe.py"
    source.write_text("VALUE = 1\n")
    fingerprint = build_deliverables.verification_fingerprint()
    (tmp_path / "docs" / "verification-results.json").write_text(json.dumps({"source_fingerprint": fingerprint}))
    assert "Отпечаток исходников совпадает" in build_deliverables.report_markdown([])
    source.write_text("VALUE = 2\n")
    assert "исторический прогон" in build_deliverables.report_markdown([])


def test_archive_ignores_hidden_secrets_editor_temp_and_symlink_targets(tmp_path, monkeypatch):
    monkeypatch.setattr(build_deliverables, "ROOT", tmp_path)
    monkeypatch.setattr(build_deliverables, "OUT", tmp_path)
    root = tmp_path / "backend"
    root.mkdir()
    (root / "app.py").write_text("pass\n")
    (root / ".env").write_text("not-for-export")
    (root / "styles.css.hoplite-write-temp").write_text("temporary")
    (root / "state.sqlite3").write_text("database")
    (root / "linked.py").symlink_to(root / "app.py")
    build_deliverables.make_source_archive()
    with zipfile.ZipFile(tmp_path / "itles_source.zip") as archive:
        assert archive.namelist() == ["itles/backend/app.py"]
    build_deliverables.make_deployment_archive()
    with zipfile.ZipFile(tmp_path / "itles_deploy.zip") as archive:
        assert archive.namelist() == ["backend/app.py"]


def test_deployment_archive_has_dockerfile_at_root_and_no_runtime_data(output_dir):
    with zipfile.ZipFile(output_dir / "itles_deploy.zip") as archive:
        names = set(archive.namelist())
        assert all((item.external_attr >> 16) & 0o777 == 0o644 for item in archive.infolist())
        assert {"Dockerfile", ".dockerignore", "deployment/compose.yaml", "server.py", "requirements.lock",
                "backend/app.py", "frontend/package-lock.json", "docs/deployment.md"} <= names
        assert names == {str(path.relative_to(build_deliverables.ROOT)) for path in build_deliverables.source_files()}
        assert not any(name.startswith(("itles/", "data/", ".local/", ".hoplite/attachments/")) for name in names)
